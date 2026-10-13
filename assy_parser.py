"""Parser sheet ASSY (Assembly Production Plan) -> tabel rapi (long format).

Format Excel:
  - baris header: kolom CUSTOMER, PART NAME, lalu kolom per tanggal (angka tanggal),
  - baris di bawahnya: "Shift 1" / "Shift 2",
  - baris "NO LOT PACKING": kode lot (mis. 30-121026-2),
  - baris data: satu baris per part, qty plan di kolom tanggal/shift,
  - kolom "Remarks" di kanan.
Bulan/tahun TIDAK ada di header tanggal, jadi ditentukan per "run" kolom (urutan tanggal
naik) dengan suara terbanyak dari kode lot. Kode lot sendiri kadang salah ketik.
"""
from __future__ import annotations

import datetime as dt
import io
import re
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from dataclasses import dataclass, field

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

MONTHS_ID = {"JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MEI": 5, "MAY": 5, "JUN": 6, "JUL": 7,
             "AGU": 8, "AGS": 8, "AUG": 8, "SEP": 9, "OKT": 10, "OCT": 10, "NOV": 11, "DES": 12, "DEC": 12}

LOT_RE = re.compile(r"^\s*\d+-(\d{2})(\d{2})(\d{2})-([12])\s*$")
SHIFT_RE = re.compile(r"^\s*shift\s*([12])\s*$", re.I)


class ParseError(Exception):
    """Format Excel tidak sesuai harapan (pesan ditampilkan ke pengguna)."""


@dataclass
class ParseResult:
    sheet: str
    period_text: str
    items: pd.DataFrame      # item_key, customer, part_name, remarks, sort_order
    plans: pd.DataFrame      # item_key, plan_date, shift, qty_plan, lot_code
    columns: pd.DataFrame    # info tiap kolom tanggal/shift + validasi total
    excluded: pd.DataFrame   # baris yang dikecualikan (stay)
    duplicates: pd.DataFrame # part kembar (customer + nama sama)
    warnings: list = field(default_factory=list)


# --------------------------------------------------------------------------- helpers
def _nk(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().upper()


def _clean(s) -> str:
    return re.sub(r"\s+", " ", str(s)).strip() if s is not None else ""


def _to_qty(v) -> int:
    if v is None or v == "" or isinstance(v, bool):
        return 0
    if isinstance(v, (int, float)):
        return int(round(v)) if v == v else 0
    s = str(v).strip()
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", s):  # "5.000" -> 5000
        s = s.replace(".", "")
    s = s.replace(",", ".")
    try:
        return int(round(float(s)))
    except ValueError:  # '#N/A', '#REF!', teks lain
        return 0


def _to_day(v):
    """Angka tanggal dari header: 16, '02', 16.0, atau datetime/date."""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and v == v:
        d = int(v)
        return d if 1 <= d <= 31 else None
    if isinstance(v, str) and v.strip().isdigit():
        d = int(v.strip())
        return d if 1 <= d <= 31 else None
    return None


def _parse_lot(v):
    m = LOT_RE.match(str(v)) if v is not None else None
    if not m:
        return None
    d, mo, y, sh = (int(x) for x in m.groups())
    if not (1 <= mo <= 12 and 1 <= d <= 31):
        return None
    return {"day": d, "month": mo, "year": 2000 + y, "shift": sh}


def _sheet_xml_info(data: bytes, sheet_name: str):
    """Kolom yang di-hide + baris akhir autofilter, dibaca langsung dari XML sheet
    (openpyxl mode read_only tidak menyediakannya, dan load penuh terlalu berat)."""
    hidden, filter_last_row = set(), None
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        rid = None
        for s in wb.find(f"{{{NS_MAIN}}}sheets"):
            if s.get("name") == sheet_name:
                rid = s.get(f"{{{NS_REL}}}id")
        if rid is None:
            return hidden, filter_last_row
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        path = None
        for r in rels:
            if r.get("Id") == rid:
                t = r.get("Target")
                path = t[1:] if t.startswith("/") else "xl/" + t
        if path is None:
            return hidden, filter_last_row
        head, tail, in_head = b"", b"", True
        with z.open(path) as f:
            while True:
                chunk = f.read(1 << 16)
                if not chunk:
                    break
                if in_head:
                    head += chunk
                    i = head.find(b"<sheetData")
                    if i != -1:
                        head, in_head = head[:i], False
                        tail = b""
                else:
                    tail = (tail + chunk)[-8192:]
    for m in re.finditer(rb"<col\s[^>]*>", head):
        attrs = dict(re.findall(rb'(\w+)="([^"]*)"', m.group(0)))
        if attrs.get(b"hidden") in (b"1", b"true"):
            lo, hi = int(attrs[b"min"]), min(int(attrs[b"max"]), 2000)
            hidden.update(range(lo, hi + 1))
    m = re.search(rb'<autoFilter[^>]*\bref="[A-Z]+\d+:[A-Z]+(\d+)"', tail)
    if m:
        filter_last_row = int(m.group(1))
    return hidden, filter_last_row


def list_sheets(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        return [s.get("name") for s in wb.find(f"{{{NS_MAIN}}}sheets")]


def _period_month(text: str):
    m = re.search(r"([A-Za-z]{3,})\.?\s*(\d{4})", text or "")
    if m and m.group(1)[:3].upper() in MONTHS_ID:
        return int(m.group(2)), MONTHS_ID[m.group(1)[:3].upper()]
    return None


# --------------------------------------------------------------------------- main
def parse_assy(data: bytes, sheet: str = "ASSY", exclude_phrases=None,
               column_mode: str = "visible", from_date: dt.date | None = None) -> ParseResult:
    """column_mode: 'visible' (kolom yang tidak di-hide) | 'from_date' | 'all'."""
    exclude_phrases = [p.strip().lower() for p in (exclude_phrases or []) if p.strip()]
    warnings: list[str] = []

    hidden_cols, filter_last_row = _sheet_xml_info(data, sheet)

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ParseError(f"Sheet '{sheet}' tidak ditemukan. Sheet yang ada: {', '.join(wb.sheetnames)}")
        ws = wb[sheet]
        rows = [tuple(r[:1000]) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    def cell(r, c):  # r, c 0-based
        if r < 0 or r >= len(rows) or c < 0 or c >= len(rows[r]):
            return None
        return rows[r][c]

    # --- header: baris yang memuat CUSTOMER dan PART NAME
    hr = col_cust = col_part = None
    for r in range(min(40, len(rows))):
        up = [_nk(v) for v in rows[r]]
        if "CUSTOMER" in up and "PART NAME" in up:
            hr, col_cust, col_part = r, up.index("CUSTOMER"), up.index("PART NAME")
            break
    if hr is None:
        raise ParseError("Baris header tidak ketemu: harus ada kolom 'CUSTOMER' dan 'PART NAME'.")

    # --- baris shift & lot
    shift_row = max(range(hr, min(hr + 4, len(rows))),
                    key=lambda r: sum(1 for v in rows[r] if v and SHIFT_RE.match(str(v))))
    if not any(v and SHIFT_RE.match(str(v)) for v in rows[shift_row]):
        raise ParseError("Baris 'Shift 1 / Shift 2' tidak ketemu di bawah header.")
    lot_row = next((r for r in range(hr, min(hr + 5, len(rows)))
                    if _nk(cell(r, col_part)) == "NO LOT PACKING"), None)
    data_start = (lot_row if lot_row is not None else shift_row) + 1

    # --- kolom Remarks
    col_rem = next((i for i, v in enumerate(rows[hr]) if _nk(v) == "REMARKS"), None)

    # --- periode (teks judul)
    period_text = ""
    for r in range(min(8, len(rows))):
        for v in rows[r]:
            if isinstance(v, str) and "period" in v.lower():
                period_text = _clean(v)
                break
        if period_text:
            break
    period_ym = _period_month(period_text)

    # --- kolom plan: ada 'Shift N' + angka tanggal
    plan_cols = []
    for c in range(len(rows[shift_row])):
        sm = SHIFT_RE.match(str(cell(shift_row, c))) if cell(shift_row, c) else None
        day = _to_day(cell(hr, c))
        if sm and day is not None:
            plan_cols.append({"c": c, "day": day, "shift": int(sm.group(1)),
                              "lot_raw": cell(lot_row, c) if lot_row is not None else None})
    if not plan_cols:
        raise ParseError("Kolom plan (tanggal + shift) tidak ketemu.")

    # --- tentukan bulan/tahun per "run" (urutan hari naik); suara terbanyak dari kode lot
    runs, cur, prev_day = [], [], None
    for pc in plan_cols:
        if isinstance(pc["day"], dt.date):
            pc["date_exact"] = pc["day"]
            pc["day"] = pc["day"].day
        d = pc["day"]
        if prev_day is not None and d < prev_day:
            runs.append(cur)
            cur = []
        cur.append(pc)
        prev_day = d
    runs.append(cur)

    prev_ym = None
    for run in runs:
        votes = Counter()
        for pc in run:
            lot = _parse_lot(pc["lot_raw"])
            pc["lot"] = lot
            if lot and lot["day"] == pc["day"]:
                votes[(lot["year"], lot["month"])] += 1
        ym = None
        if votes:
            top = votes.most_common(2)
            if len(top) == 1 or top[0][1] > top[1][1]:
                ym = top[0][0]
        if ym is None and prev_ym:
            y, m = prev_ym
            ym = (y + (m == 12), 1 if m == 12 else m + 1)
        if ym is None and period_ym:
            ym = period_ym
        if ym is None:
            raise ParseError("Bulan/tahun kolom tanggal tidak bisa ditentukan (kode lot tidak terbaca).")
        if prev_ym and ym <= prev_ym:
            warnings.append(f"Urutan bulan tidak naik di kolom {get_column_letter(run[0]['c'] + 1)} "
                            f"({ym[0]}-{ym[1]:02d} setelah {prev_ym[0]}-{prev_ym[1]:02d}). Cek header tanggal.")
        prev_ym = ym
        for pc in run:
            pc["ym"] = ym

    last_run_ym = runs[-1][0]["ym"]
    if period_ym and last_run_ym != period_ym:
        warnings.append(f"Bulan kolom terakhir ({last_run_ym[0]}-{last_run_ym[1]:02d}) berbeda dengan "
                        f"judul sheet ({period_text}). Cek apakah file ini benar.")

    for pc in plan_cols:
        y, m = pc["ym"]
        try:
            pc["date"] = pc.get("date_exact") or dt.date(y, m, pc["day"])
        except ValueError:
            pc["date"] = None
            warnings.append(f"Kolom {get_column_letter(pc['c'] + 1)}: tanggal {pc['day']} tidak ada di bulan {y}-{m:02d}, dilewati.")
        lot = pc["lot"]
        pc["lot_ok"] = bool(lot and pc["date"] and (lot["year"], lot["month"], lot["day"]) == (pc["date"].year, pc["date"].month, pc["date"].day)
                            and lot["shift"] == pc["shift"])
        pc["hidden"] = (pc["c"] + 1) in hidden_cols

    # --- pilih kolom yang dibaca
    for pc in plan_cols:
        if pc["date"] is None:
            pc["selected"] = False
        elif column_mode == "visible":
            pc["selected"] = not pc["hidden"]
        elif column_mode == "from_date":
            pc["selected"] = bool(from_date and pc["date"] >= from_date)
        else:
            pc["selected"] = True
    sel = [pc for pc in plan_cols if pc["selected"]]
    if not sel:
        warnings.append("Tidak ada kolom tanggal yang terpilih. Ubah pilihan 'Kolom yang dibaca'.")

    # --- baris total (SUBTOTAL) di atas header
    tot_row = hr - 1
    # --- baca baris data
    end_row = filter_last_row if filter_last_row else len(rows)  # 1-based inklusif
    end_row = min(end_row, len(rows))
    occ = Counter()
    items, plans, excluded = [], [], []
    parsed_tot, excl_tot = Counter(), Counter()
    for r in range(data_start, end_row):
        part = _clean(cell(r, col_part))
        if not part:
            continue
        if re.match(r"^(total|ttl)\b", part, re.I):
            break
        cust = _clean(cell(r, col_cust))
        qtys = {pc["c"]: _to_qty(cell(r, pc["c"])) for pc in sel}
        if any(p in part.lower() for p in exclude_phrases):
            excluded.append({"row": r + 1, "customer": cust, "part_name": part,
                             "total_qty": sum(qtys.values())})
            for c, q in qtys.items():
                excl_tot[c] += q
            continue
        base = f"{_nk(cust) or '-'}|{_nk(part)}"
        occ[base] += 1
        key = base if occ[base] == 1 else f"{base}|#{occ[base]}"
        for c, q in qtys.items():
            parsed_tot[c] += q
        if any(q > 0 for q in qtys.values()):
            items.append({"item_key": key, "customer": cust, "part_name": part,
                          "remarks": _clean(cell(r, col_rem)) if col_rem is not None else "",
                          "sort_order": r + 1, "base": base, "occ": occ[base]})
            for pc in sel:
                q = qtys[pc["c"]]
                if q > 0:
                    plans.append({"item_key": key, "plan_date": pc["date"].isoformat(),
                                  "shift": pc["shift"], "qty_plan": q,
                                  "lot_code": _clean(pc["lot_raw"])})

    # --- tabel kolom + validasi total
    col_rows = []
    for pc in plan_cols:
        c = pc["c"]
        excel_total = _to_qty(cell(tot_row, c)) if tot_row >= 0 else None
        p, e = parsed_tot.get(c, 0), excl_tot.get(c, 0)
        status = ""
        if pc["selected"]:
            if excel_total == p:
                status = "OK"
            elif excel_total == p + e:
                status = "OK (selisih = baris stay)"
            else:
                status = "CEK"
        col_rows.append({"kolom": get_column_letter(c + 1), "tanggal": pc["date"], "shift": pc["shift"],
                         "kode_lot": _clean(pc["lot_raw"]), "lot_cocok": pc["lot_ok"],
                         "di_hide": pc["hidden"], "dibaca": pc["selected"],
                         "total_excel": excel_total if pc["selected"] else None,
                         "total_dibaca": p if pc["selected"] else None,
                         "total_stay_dibuang": e if pc["selected"] else None, "status": status})
    columns_df = pd.DataFrame(col_rows)

    bad_lot = [r for r in col_rows if r["dibaca"] and not r["lot_cocok"]]
    if bad_lot:
        warnings.append("Kode lot tidak cocok dengan tanggal/shift di kolom: "
                        + ", ".join(f"{r['kolom']} ({r['kode_lot'] or 'kosong'})" for r in bad_lot[:12])
                        + (" ..." if len(bad_lot) > 12 else "")
                        + ". Tanggal tetap diambil dari header + urutan kolom, bukan dari kode lot.")
    if (columns_df["status"] == "CEK").any():
        warnings.append("Total hasil baca berbeda dengan baris 'Plan' di Excel pada kolom: "
                        + ", ".join(columns_df.loc[columns_df["status"] == "CEK", "kolom"])
                        + ". Bisa karena filter aktif di Excel atau ada baris yang salah terbaca.")

    items_df = pd.DataFrame(items, columns=["item_key", "customer", "part_name", "remarks", "sort_order", "base", "occ"])
    dup_bases = items_df.loc[items_df["occ"] > 1, "base"].unique() if len(items_df) else []
    # part kembar: tandai juga baris pertamanya
    dup_df = items_df[items_df["base"].isin(dup_bases)].drop(columns=["base", "occ"]) if len(items_df) else items_df
    if len(dup_df):
        warnings.append(f"Ada {len(dup_bases)} part kembar (customer + nama sama). Dibedakan lewat urutan baris; "
                        "kolom Remarks ditampilkan di samping nama part.")
    return ParseResult(
        sheet=sheet, period_text=period_text,
        items=items_df.drop(columns=["base", "occ"]),
        plans=pd.DataFrame(plans, columns=["item_key", "plan_date", "shift", "qty_plan", "lot_code"]),
        columns=columns_df, excluded=pd.DataFrame(excluded, columns=["row", "customer", "part_name", "total_qty"]),
        duplicates=dup_df, warnings=warnings)
